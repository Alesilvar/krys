from __future__ import annotations

import csv
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from extractor import ExtractionError

from .forms import ManualRecordForm, PresentationEditForm, UploadLettersForm
from .models import AuditEvent, Equipment, Evaluation, Letter, Presentation
from .services import persist_preview, prepare_preview


def records_queryset(request):
    queryset = Presentation.objects.select_related(
        "cpc_letter", "equipment", "evaluation", "evaluation__response_letter"
    )
    query = request.GET.get("q", "").strip()
    result = request.GET.get("result", "").strip()
    if query:
        queryset = queryset.filter(
            Q(equipment__code__icontains=query)
            | Q(cpc_letter__full_number__icontains=query)
            | Q(equipment_type__icontains=query)
            | Q(evaluation__response_letter__full_number__icontains=query)
        )
    if result:
        queryset = queryset.filter(evaluation__result=result)
    return queryset, query, result


@login_required
def dashboard(request):
    records, query, result = records_queryset(request)
    context = {
        "records": records,
        "query": query,
        "selected_result": result,
        "total": Presentation.objects.count(),
        "approved": Evaluation.objects.filter(result=Evaluation.Result.APPROVED).count(),
        "observed": Evaluation.objects.filter(result=Evaluation.Result.OBSERVED).count(),
        "pending_letters": Letter.objects.filter(
            direction=Letter.Direction.CPC_AAM, response_letter__isnull=True
        ).count(),
    }
    return render(request, "records/dashboard.html", context)


@login_required
def upload_letters(request):
    preview = request.session.get("upload_preview")
    if request.method == "POST":
        form = UploadLettersForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                preview = prepare_preview(form.cleaned_data["files"])
                request.session["upload_preview"] = preview
                request.session.modified = True
                return render(
                    request,
                    "records/upload.html",
                    {"form": UploadLettersForm(), "preview": preview},
                )
            except (ExtractionError, IntegrityError, ValueError) as error:
                messages.error(request, str(error))
    else:
        form = UploadLettersForm()
    return render(request, "records/upload.html", {"form": form, "preview": preview})


@login_required
def confirm_upload(request):
    if request.method != "POST":
        return redirect("upload_letters")
    preview = request.session.get("upload_preview")
    if not preview:
        messages.error(request, "La vista previa expiró. Vuelve a cargar los archivos.")
        return redirect("upload_letters")
    try:
        persist_preview(preview, request.POST, request.user)
    except (ExtractionError, IntegrityError, ValidationError, ValueError) as error:
        messages.error(request, str(error))
        return render(
            request,
            "records/upload.html",
            {"form": UploadLettersForm(), "preview": preview},
        )
    request.session.pop("upload_preview", None)
    messages.success(request, "Las cartas y sus equipos se guardaron correctamente.")
    return redirect("dashboard")


def _correlative(full_number):
    return int(full_number.rsplit("-", 1)[-1])


@login_required
def manual_record(request):
    form = ManualRecordForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        data = form.cleaned_data
        try:
            with transaction.atomic():
                cpc = Letter.objects.filter(full_number=data["cpc_full_number"]).first()
                if cpc:
                    if cpc.direction != Letter.Direction.CPC_AAM:
                        raise ValidationError("El número indicado no corresponde a una carta CPC.")
                    if cpc.letter_date != data["cpc_date"]:
                        raise ValidationError("La fecha CPC no coincide con la carta ya registrada.")
                    if cpc.reception_date != data["reception_date"]:
                        raise ValidationError("La fecha de recepción no coincide con la carta ya registrada.")
                else:
                    predecessor = None
                    if data["predecessor_aam"]:
                        predecessor = Letter.objects.filter(
                            full_number=data["predecessor_aam"],
                            direction=Letter.Direction.AAM_CPC,
                        ).first()
                        if predecessor is None:
                            raise ValidationError("La respuesta AAM anterior aún no está registrada.")
                    cpc = Letter.objects.create(
                        full_number=data["cpc_full_number"],
                        correlative=_correlative(data["cpc_full_number"]),
                        direction=Letter.Direction.CPC_AAM,
                        letter_date=data["cpc_date"],
                        reception_date=data["reception_date"],
                        predecessor_response=predecessor,
                        source=Letter.Source.MANUAL,
                        pdf=data["cpc_pdf"],
                        created_by=request.user,
                    )
                if data["cpc_pdf"] and not cpc.pdf:
                    cpc.pdf = data["cpc_pdf"]
                    cpc.save(update_fields=["pdf", "updated_at"])

                response = getattr(cpc, "response_letter", None)
                if data["aam_full_number"]:
                    by_number = Letter.objects.filter(full_number=data["aam_full_number"]).first()
                    if response and response.full_number != data["aam_full_number"]:
                        raise ValidationError("Esta carta CPC ya tiene otra respuesta AAM.")
                    if by_number and by_number.response_to_id != cpc.id:
                        raise ValidationError("La respuesta AAM ya está vinculada a otra carta CPC.")
                    response = response or by_number
                    if response and response.letter_date != data["aam_date"]:
                        raise ValidationError("La fecha AAM no coincide con la carta ya registrada.")
                    if response is None:
                        response = Letter.objects.create(
                            full_number=data["aam_full_number"],
                            correlative=_correlative(data["aam_full_number"]),
                            direction=Letter.Direction.AAM_CPC,
                            letter_date=data["aam_date"],
                            source=Letter.Source.MANUAL,
                            pdf=data["aam_pdf"],
                            response_to=cpc,
                            created_by=request.user,
                        )
                    elif data["aam_pdf"] and not response.pdf:
                        response.pdf = data["aam_pdf"]
                        response.save(update_fields=["pdf", "updated_at"])

                equipment, _ = Equipment.objects.get_or_create(
                    code=data["equipment_code"].strip().upper()
                )
                presentation = Presentation.objects.create(
                    cpc_letter=cpc,
                    equipment=equipment,
                    equipment_type=data["equipment_type"].strip(),
                    additional_data=data["additional_data"].strip(),
                    created_by=request.user,
                )
                if data["result"]:
                    Evaluation.objects.create(
                        presentation=presentation,
                        response_letter=response,
                        result=data["result"],
                        detail=data["detail"].strip(),
                        updated_by=request.user,
                    )
                AuditEvent.objects.create(
                    user=request.user,
                    action=AuditEvent.Action.CREATE,
                    object_type="Presentation",
                    object_id=str(presentation.pk),
                    details={"source": "manual"},
                )
        except (IntegrityError, ValidationError) as error:
            messages.error(request, str(error))
        else:
            messages.success(request, "Registro manual guardado.")
            return redirect("dashboard")
    return render(request, "records/manual.html", {"form": form})


@login_required
def edit_presentation(request, pk):
    presentation = get_object_or_404(
        Presentation.objects.select_related(
            "cpc_letter", "equipment", "evaluation", "evaluation__response_letter"
        ),
        pk=pk,
    )
    form = PresentationEditForm.from_presentation(presentation, request.POST or None)
    if request.method == "POST" and form.is_valid():
        old = {
            "cpc_full_number": presentation.cpc_letter.full_number,
            "cpc_date": str(presentation.cpc_letter.letter_date),
            "reception_date": str(presentation.cpc_letter.reception_date),
            "equipment_type": presentation.equipment_type,
            "equipment_code": presentation.equipment.code,
            "additional_data": presentation.additional_data,
        }
        try:
            with transaction.atomic():
                equipment, _ = Equipment.objects.get_or_create(
                    code=form.cleaned_data["equipment_code"].strip().upper()
                )
                cpc = presentation.cpc_letter
                cpc.full_number = form.cleaned_data["cpc_full_number"]
                cpc.correlative = _correlative(cpc.full_number)
                cpc.letter_date = form.cleaned_data["cpc_date"]
                cpc.reception_date = form.cleaned_data["reception_date"]
                if form.cleaned_data["cpc_pdf"]:
                    cpc.pdf = form.cleaned_data["cpc_pdf"]
                cpc.save()
                presentation.equipment = equipment
                presentation.equipment_type = form.cleaned_data["equipment_type"].strip()
                presentation.additional_data = form.cleaned_data["additional_data"].strip()
                presentation.save()

                result = form.cleaned_data["result"]
                response = getattr(presentation.cpc_letter, "response_letter", None)
                if form.cleaned_data["aam_full_number"]:
                    if response is None:
                        response = Letter.objects.create(
                            full_number=form.cleaned_data["aam_full_number"],
                            correlative=_correlative(form.cleaned_data["aam_full_number"]),
                            direction=Letter.Direction.AAM_CPC,
                            letter_date=form.cleaned_data["aam_date"],
                            source=Letter.Source.MANUAL,
                            pdf=form.cleaned_data["aam_pdf"],
                            response_to=cpc,
                            created_by=request.user,
                        )
                    else:
                        response.full_number = form.cleaned_data["aam_full_number"]
                        response.correlative = _correlative(response.full_number)
                        response.letter_date = form.cleaned_data["aam_date"]
                        if form.cleaned_data["aam_pdf"]:
                            response.pdf = form.cleaned_data["aam_pdf"]
                        response.save()
                if result:
                    Evaluation.objects.update_or_create(
                        presentation=presentation,
                        defaults={
                            "response_letter": response,
                            "result": result,
                            "detail": form.cleaned_data["detail"].strip(),
                            "updated_by": request.user,
                        },
                    )
                else:
                    Evaluation.objects.filter(presentation=presentation).delete()
                AuditEvent.objects.create(
                    user=request.user,
                    action=AuditEvent.Action.UPDATE,
                    object_type="Presentation",
                    object_id=str(presentation.pk),
                    details={
                        "before": old,
                        "after": {
                            key: str(value) if value is not None else ""
                            for key, value in form.cleaned_data.items()
                        },
                    },
                )
        except (IntegrityError, ValidationError) as error:
            messages.error(request, str(error))
        else:
            messages.success(request, "Registro actualizado.")
            return redirect("dashboard")
    return render(
        request,
        "records/edit.html",
        {"form": form, "presentation": presentation},
    )


@login_required
def delete_letter(request, pk):
    letter = get_object_or_404(Letter, pk=pk, direction=Letter.Direction.CPC_AAM)
    if request.method == "POST":
        response = getattr(letter, "response_letter", None)
        files = [item.pdf.name for item in (letter, response) if item and item.pdf]
        number = letter.full_number
        with transaction.atomic():
            AuditEvent.objects.create(
                user=request.user,
                action=AuditEvent.Action.DELETE,
                object_type="Letter",
                object_id=number,
                details={"permanent": True},
            )
            letter.delete()
        from django.core.files.storage import default_storage

        for name in files:
            default_storage.delete(name)
        messages.success(request, f"La carta {number} y todos sus datos fueron eliminados.")
        return redirect("dashboard")
    return render(request, "records/delete.html", {"letter": letter})


def export_values(presentation):
    try:
        evaluation = presentation.evaluation
        result = evaluation.get_result_display()
        detail = evaluation.detail
        response = evaluation.response_letter
    except Evaluation.DoesNotExist:
        result = detail = ""
        response = None
    return [
        presentation.owner,
        presentation.cpc_letter.correlative,
        presentation.cpc_letter.letter_date.strftime("%d/%m/%Y"),
        presentation.cpc_letter.reception_date.strftime("%d/%m/%Y")
        if presentation.cpc_letter.reception_date
        else "",
        presentation.procedure_type,
        presentation.equipment_type,
        presentation.equipment.code,
        result,
        detail,
        response.correlative if response else "",
        response.letter_date.strftime("%d/%m/%Y") if response else "",
        presentation.additional_data,
    ]


EXPORT_HEADERS = [
    "Propietario del Equipo",
    "N° Correlativo Carta de Entrada",
    "Fecha de carta CPC",
    "Fecha de recepción de carta",
    "Tipo de trámite",
    "Tipo de Equipo",
    "Placa / Código del Equipo",
    "Resultado de la Evaluación SST",
    "Detalle de la Observación o Conformidad",
    "N° Correlativo Carta Respuesta AAM",
    "Fecha de carta AAM",
    "Dato adicional",
]


@login_required
def export_csv(request):
    records, _, _ = records_queryset(request)
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="registros-krys.csv"'
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow(EXPORT_HEADERS)
    for record in records:
        writer.writerow(export_values(record))
    return response


@login_required
def export_xlsx(request):
    records, _, _ = records_queryset(request)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registros"
    sheet.append(EXPORT_HEADERS)
    for record in records:
        sheet.append(export_values(record))
    fill = PatternFill("solid", fgColor="163A5F")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(EXPORT_HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(header) + 2, 14), 42
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="registros-krys.xlsx"'
    return response
