"""Transformacion de registros a formatos de exportacion."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Evaluation


EXPORT_HEADERS = [
    "Propietario del Equipo",
    "N° Correlativo Carta de Entrada",
    "Fecha de carta CPC",
    "Fecha de recepción de carta",
    "Tipo de trámite",
    "Tipo de Equipo",
    "Placa / Código del Equipo",
    "Resultado de la Evaluación SST",
    "Detalle de la Observación o Conformidad",
    "N° Correlativo Carta Respuesta AAM",
    "Fecha de carta AAM",
    "Dato adicional",
]


def presentation_values(presentation):
    """Convierte una presentacion en una fila neutral para CSV o Excel."""
    try:
        evaluation = presentation.evaluation
        result = evaluation.get_result_display()
        detail = evaluation.detail
        response = evaluation.response_letter
    except Evaluation.DoesNotExist:
        result = detail = ""
        response = None
    return [
        presentation.owner,
        presentation.cpc_letter.correlative,
        presentation.cpc_letter.letter_date.strftime("%d/%m/%Y"),
        presentation.cpc_letter.reception_date.strftime("%d/%m/%Y")
        if presentation.cpc_letter.reception_date
        else "",
        presentation.procedure_type,
        presentation.equipment_type,
        presentation.equipment.code,
        result,
        detail,
        response.correlative if response else "",
        response.letter_date.strftime("%d/%m/%Y") if response else "",
        presentation.additional_data,
    ]


def workbook_bytes(records):
    """Genera el contenido binario de un libro Excel con los registros."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registros"
    sheet.append(EXPORT_HEADERS)
    for record in records:
        sheet.append(presentation_values(record))

    fill = PatternFill("solid", fgColor="163A5F")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(EXPORT_HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(header) + 2, 14), 42
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
